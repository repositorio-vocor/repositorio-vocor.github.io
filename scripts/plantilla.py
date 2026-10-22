"""Genera plantilla/vocor-planilla-modelo.xlsx. Correr de nuevo si cambian los vocabularios."""
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).parent))
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from vocor_datos import AREAS, FORMATOS, PUBLICOS, RAIZ, cargar_materiales

wb = Workbook()
ws = wb.active; ws.title = "Materiales"
cols = [("Recurso", 42), ("UE", 24), ("Área", 36), ("Formato", 24), ("Público", 30), ("Palabras clave", 30), ("Descripción", 60), ("Links", 42), ("Imagen", 36)]
ws.append([c for c, _ in cols])
ws.append(["Ejemplo: Plantas nativas de las Sierras (borrá esta fila)", "IMBIV; IDEA", "Ciencias Biológicas y de la Salud", "Podcast", "Público General; Docentes",
           "plantas nativas, flora, sierras", "Podcast sobre la flora nativa de Córdoba.", "https://open.spotify.com/show/…", "https://ejemplo.org/portada.jpg"])
fill = PatternFill("solid", fgColor="1F6B52")
for i, (c, w) in enumerate(cols, start=1):
    cell = ws.cell(row=1, column=i); cell.font = Font(bold=True, color="FFFFFF"); cell.fill = fill
    ws.column_dimensions[cell.column_letter].width = w
for c in ws[2]: c.font = Font(italic=True, color="777777"); c.alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = "A2"

v = wb.create_sheet("Valores")
ues = sorted({u for m in cargar_materiales() for u in m.get("ue", [])})
v.append(["Área", "Público", "Formato", "UE ya cargadas"])
for c in v[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = fill
for i in range(max(len(AREAS), len(PUBLICOS), len(FORMATOS), len(ues))):
    v.append([(AREAS + [""] * 50)[i], (PUBLICOS + [""] * 50)[i], (FORMATOS + [""] * 50)[i], (ues + [""] * 50)[i]])
for col, w in zip("ABCD", (42, 30, 38, 46)): v.column_dimensions[col].width = w

h = wb.create_sheet("Instrucciones")
for linea in [
    "Cómo completar la planilla",
    "",
    "1. Una fila por material. Borrá la fila de ejemplo (si la dejás, se ignora).",
    "2. Obligatorios: Recurso (título), UE y Formato.",
    "3. Si una celda tiene varios valores, separalos con punto y coma ( ; ). Ej.: IDEA; IMBIV",
    "4. Área, Público y Formato: usá los valores de la hoja Valores.",
    "5. Links: pegá la dirección completa (https://…). Si hay varias, separalas con un espacio.",
    "6. Imagen (opcional): dirección web de la portada. Se descarga sola. También podés agregarla después desde /admin.",
    "7. Guardá como .xlsx y subila a la carpeta importar/ del repositorio en GitHub (botón Add file → Upload files).",
    "8. En unos minutos los materiales aparecen en el sitio. El reporte de qué filas entraron queda en importar/procesadas/.",
    "",
    "Si un material ya existe con el mismo título, no se duplica.",
]:
    h.append([linea])
h["A1"].font = Font(bold=True, size=14); h.column_dimensions["A"].width = 110
wb.save(RAIZ / "plantilla" / "vocor-planilla-modelo.xlsx")
print("ok")
