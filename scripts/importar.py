"""Importa las planillas que haya en importar/ y las convierte en materiales.

Lo ejecuta GitHub Actions cada vez que alguien sube una planilla a la carpeta
importar/. También se puede correr a mano: python scripts/importar.py

Por cada planilla:
  - crea un archivo en content/materiales/ por cada fila válida;
  - si la fila trae una URL en la columna Imagen, descarga la portada;
  - mueve la planilla a importar/procesadas/ junto con un reporte (.md)
    que dice qué filas entraron y cuáles no, y por qué.
"""
import csv
import io
import os
import shutil
import sys
import urllib.request
from datetime import datetime, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from vocor_datos import (AREAS, IMAGENES, PUBLICOS, RAIZ, CONTENIDO, cargar_materiales, guardar_material,
                         nombre_libre, norm, separar, separar_links, separar_palabras, slugify)

CARPETA = RAIZ / "importar"
PROCESADAS = CARPETA / "procesadas"

COLUMNAS = {
    "recurso": ["recurso", "titulo", "título", "nombre"],
    "ue": ["ue", "unidad ejecutora", "instituto", "unidad", "unidades ejecutoras"],
    "area": ["area", "área", "areas", "áreas"],
    "formato": ["formato", "formatos", "tipo"],
    "publico": ["publico", "público", "destinatarios"],
    "palabras": ["palabras clave", "palabras claves", "keywords", "temas"],
    "descripcion": ["descripcion", "descripción"],
    "links": ["links", "link", "enlace", "enlaces", "url"],
    "imagen": ["imagen", "portada", "imagen (url)", "url de la imagen"],
}


def leer_filas(ruta):
    """Devuelve (encabezados, filas) de la primera hoja de un .xlsx o de un .csv."""
    if ruta.suffix.lower() == ".csv":
        texto = ruta.read_bytes()
        for cod in ("utf-8-sig", "latin-1"):
            try:
                texto = texto.decode(cod)
                break
            except UnicodeDecodeError:
                continue
        encabezado = texto.splitlines()[0] if texto.strip() else ""
        separador = max(";,\t", key=encabezado.count)  # Excel en español guarda los CSV con ";"
        filas = list(csv.reader(io.StringIO(texto), delimiter=separador))
    else:
        from openpyxl import load_workbook
        wb = load_workbook(ruta, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        filas = [["" if c is None else str(c) for c in fila] for fila in ws.iter_rows(values_only=True)]
    filas = [f for f in filas if any(str(c).strip() for c in f)]
    if not filas:
        return [], []
    return [norm(h) for h in filas[0]], filas[1:]


def mapear(encabezados):
    idx = {}
    for campo, alias in COLUMNAS.items():
        for i, h in enumerate(encabezados):
            if h in [norm(a) for a in alias]:
                idx[campo] = i
                break
    return idx


def bajar_imagen(url, slug):
    from PIL import Image
    req = urllib.request.Request(url, headers={"User-Agent": "VoCor-importador/1.0"})
    with urllib.request.urlopen(req, timeout=20) as r:
        datos = r.read(15 * 1024 * 1024)
    im = Image.open(io.BytesIO(datos)).convert("RGB")
    im.thumbnail((1200, 1200))
    IMAGENES.mkdir(exist_ok=True)
    destino = IMAGENES / f"{slug}.jpg"
    im.save(destino, "JPEG", quality=82, optimize=True, progressive=True)
    return f"/imagenes/{slug}.jpg"


def importar_planilla(ruta, existentes_titulos, slugs):
    encabezados, filas = leer_filas(ruta)
    idx = mapear(encabezados)
    reporte = []
    if "recurso" not in idx:
        reporte.append(("—", "Error", "", "No encontré la columna Recurso. Usá la planilla modelo."))
        return 0, reporte
    ahora = datetime.now(timezone.utc)
    nuevos = 0
    for n, fila in enumerate(filas, start=2):
        celda = lambda c: str(fila[idx[c]]).strip() if c in idx and idx[c] < len(fila) else ""
        d = {
            "recurso": celda("recurso"),
            "ue": separar(celda("ue"), "ue"),
            "area": separar(celda("area"), "area"),
            "formato": separar(celda("formato"), "formato"),
            "publico": separar(celda("publico"), "publico"),
            "palabras": separar_palabras(celda("palabras")),
            "descripcion": celda("descripcion"),
            "links": separar_links(celda("links")),
        }
        if d["recurso"].lower().startswith("ejemplo:"):
            reporte.append((n, "Omitida", d["recurso"], "Es la fila de ejemplo de la planilla modelo."))
            continue
        faltan = [nom for nom, ok in (("Recurso", d["recurso"]), ("UE", d["ue"]), ("Formato", d["formato"])) if not ok]
        if faltan:
            reporte.append((n, "Error", d["recurso"], "Falta: " + ", ".join(faltan)))
            continue
        if norm(d["recurso"]) in existentes_titulos:
            reporte.append((n, "Duplicado", d["recurso"], "Ya hay un material con ese título."))
            continue
        avisos = []
        nuevos_valores = [v for v in d["area"] if v not in AREAS] + [v for v in d["publico"] if v not in PUBLICOS]
        if nuevos_valores:
            avisos.append("Valores fuera de la lista: " + ", ".join(nuevos_valores))
        if not d["links"]:
            avisos.append("Sin enlace")
        slug = nombre_libre(slugify(d["recurso"]), slugs)
        url_img = celda("imagen")
        if url_img.startswith("http"):
            try:
                d["imagen"] = bajar_imagen(url_img, slug)
            except Exception as e:  # noqa: BLE001
                avisos.append(f"No se pudo descargar la imagen ({type(e).__name__})")
        d["origen"] = "excel"
        d["creado"] = ahora.replace(microsecond=0).isoformat().replace("+00:00", "Z")
        guardar_material(slug, d)
        existentes_titulos.add(norm(d["recurso"]))
        nuevos += 1
        reporte.append((n, "Importada" if not avisos else "Importada con avisos", d["recurso"], "; ".join(avisos)))
    return nuevos, reporte


def main():
    planillas = sorted(p for p in CARPETA.glob("*") if p.is_file() and p.suffix.lower() in (".xlsx", ".csv") and not p.name.startswith("~$"))
    if not planillas:
        print("No hay planillas para importar.")
        return 0
    existentes = cargar_materiales()
    titulos = {norm(m.get("recurso")) for m in existentes}
    slugs = {m["id"] for m in existentes}
    PROCESADAS.mkdir(parents=True, exist_ok=True)
    resumen = ["## Importación de planillas", ""]
    for ruta in planillas:
        sello = datetime.now().strftime("%Y-%m-%d-%H%M")
        try:
            nuevos, reporte = importar_planilla(ruta, titulos, slugs)
        except Exception as e:  # noqa: BLE001
            nuevos, reporte = 0, [("—", "Error", "", f"No se pudo leer el archivo: {e}")]
        destino = PROCESADAS / f"{sello}-{ruta.name}"
        shutil.move(str(ruta), destino)
        lineas = [f"# Reporte de importación: {ruta.name}", "", f"Procesada: {sello}", f"Materiales nuevos: **{nuevos}**", "",
                  "| Fila | Estado | Recurso | Detalle |", "|---|---|---|---|"]
        lineas += [f"| {f} | {e} | {r.replace('|', '/')} | {det.replace('|', '/')} |" for f, e, r, det in reporte]
        (PROCESADAS / f"{sello}-{ruta.stem}-reporte.md").write_text("\n".join(lineas) + "\n", encoding="utf-8")
        resumen += [f"### {ruta.name}: {nuevos} materiales nuevos", ""] + lineas[5:] + [""]
        print(f"{ruta.name}: {nuevos} materiales nuevos")
    if os.environ.get("GITHUB_STEP_SUMMARY"):
        with open(os.environ["GITHUB_STEP_SUMMARY"], "a", encoding="utf-8") as f:
            f.write("\n".join(resumen) + "\n")
    return 0


if __name__ == "__main__":
    sys.exit(main())
